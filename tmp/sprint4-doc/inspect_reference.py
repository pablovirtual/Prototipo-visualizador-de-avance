from pathlib import Path
from docx import Document


path = Path(r"output/docx/Actividad_Integradora_2_Avance_actual.docx").resolve()
doc = Document(path)

print(f"PARAGRAPHS={len(doc.paragraphs)} TABLES={len(doc.tables)} SECTIONS={len(doc.sections)}")
for index, paragraph in enumerate(doc.paragraphs):
    text = paragraph.text.strip()
    if text:
        print(f"P{index:03d} [{paragraph.style.name}] {text}")

for table_index, table in enumerate(doc.tables):
    print(f"TABLE {table_index}: {len(table.rows)}x{len(table.columns)}")
    for row in table.rows:
        print(" | ".join(cell.text.replace("\n", " / ") for cell in row.cells))

for section_index, section in enumerate(doc.sections):
    print(f"HEADER {section_index}: " + " || ".join(p.text for p in section.header.paragraphs))
    print(f"FOOTER {section_index}: " + " || ".join(p.text for p in section.footer.paragraphs))

props = doc.core_properties
print(f"CORE author={props.author!r} title={props.title!r} subject={props.subject!r}")
