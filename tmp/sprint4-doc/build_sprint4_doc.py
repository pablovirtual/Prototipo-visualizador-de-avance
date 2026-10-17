from copy import deepcopy
from pathlib import Path
import shutil

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
REFERENCE = ROOT / "output" / "docx" / "Actividad_Integradora_2_Avance_actual.docx"
FINAL = ROOT / "output" / "docx" / "Actividad_3.2_Cuarto_sprint.docx"

BLACK = "000000"
DARK_BLUE = "17365D"
PALE_BLUE = "EAF2F8"
LIGHT_GRAY = "D9D9D9"


def set_run_font(run, size=12, bold=None, color=BLACK, italic=None):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "6")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), LIGHT_GRAY)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_hyperlink(paragraph, text, url, bold_lead=None):
    if bold_lead:
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, bold=True)
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "24")
    r_pr.extend([fonts, color, underline, size])
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_paragraph(doc, text="", style=None, align=None, before=0, after=6, line=1.5):
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    if align is not None:
        paragraph.alignment = align
    if text:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(0 if level == 1 else 6)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    set_run_font(run, size=14 if level == 1 else 12, bold=True)
    return paragraph


def add_bullet(doc, text):
    paragraph = add_paragraph(doc, style="List Bullet", after=3, line=1.25)
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)

    for index, (cell, header, width) in enumerate(zip(table.rows[0].cells, headers, widths)):
        cell.width = Inches(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        set_cell_shading(cell, DARK_BLUE)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        run = paragraph.add_run(header)
        set_run_font(run, size=9.5, bold=True, color="FFFFFF")
    header_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_pr.append(repeat)
    prevent_row_split(table.rows[0])

    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for column_index, (cell, value, width) in enumerate(zip(cells, values, widths)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if column_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            )
            run = paragraph.add_run(value)
            set_run_font(run, size=9.3, bold=(column_index == 0))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def page_break(doc):
    doc.add_page_break()


def add_page_number_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "2"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])
    set_run_font(run, size=9)


shutil.copy2(REFERENCE, FINAL)
doc = Document(FINAL)

# Preserve the section properties and template relationships while replacing the body.
body = doc._element.body
sect_pr = body.sectPr
for child in list(body):
    if child is not sect_pr:
        body.remove(child)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
normal.font.size = Pt(12)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(6)

for style_name, size in (("Title", 20), ("Heading 1", 14), ("Heading 2", 12)):
    style = doc.styles[style_name]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)

props = doc.core_properties
props.author = "Pedro Pablo Rodríguez Gómez"
props.last_modified_by = "Pedro Pablo Rodríguez Gómez"
props.title = "Actividad 3.2 Cuarto sprint"
props.subject = "Resultado del desarrollo y retrospectiva del Sprint 4"
props.keywords = "Scrum, Sprint 4, Vue, filtros, avance curricular"

section = doc.sections[0]
section.different_first_page_header_footer = True
footer = section.footer
footer_paragraph = footer.paragraphs[0]
footer_paragraph.clear()
footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_paragraph.paragraph_format.space_after = Pt(0)
footer_run = footer_paragraph.add_run("Actividad 3.2 Cuarto sprint  |  Página ")
set_run_font(footer_run, size=9, color="666666")
add_page_number_field(footer_paragraph)

first_footer = section.first_page_footer
for paragraph in first_footer.paragraphs:
    paragraph.clear()

# Página 1. Portada
add_paragraph(doc, "", after=12)
p = add_paragraph(doc, "UNIVERSIDAD DE GUADALAJARA", align=WD_ALIGN_PARAGRAPH.CENTER, after=34)
for run in p.runs:
    set_run_font(run, size=15, bold=True)

p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(22)
r = p.add_run("Actividad 3.2\nCuarto sprint")
set_run_font(r, size=20, bold=True)

p = add_paragraph(
    doc,
    "Visualizador de Avance Curricular y Disponibilidad de Materias por Prerrequisitos",
    align=WD_ALIGN_PARAGRAPH.CENTER,
    after=52,
)
for run in p.runs:
    set_run_font(run, size=14, bold=True)

for text in (
    "Alumno: Pedro Pablo Rodríguez Gómez",
    "Materia: Proyecto VII",
    "Asesora: María Luisa García Sánchez",
    "Fecha: 16 de octubre de 2026",
):
    p = add_paragraph(doc, text, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    for run in p.runs:
        set_run_font(run, size=12)

page_break(doc)

# Página 2. Introducción y planificación
add_heading(doc, "1. Introducción")
add_paragraph(
    doc,
    "El cuarto sprint dio continuidad al prototipo del Visualizador de Avance Curricular mediante una función de consulta que facilita localizar materias de acuerdo con su estado académico. El incremento se construyó sobre la gestión de prerrequisitos incorporada en el sprint anterior y conserva la malla de prueba compuesta por 12 materias y 40 créditos.",
)
add_paragraph(
    doc,
    "Al concluir el ciclo, la aplicación permite mostrar todas las asignaturas o concentrarse en las cursadas, disponibles y bloqueadas. La solución actualiza los resultados en respuesta al avance registrado, mantiene la persistencia local y se encuentra integrada en la rama main y publicada en GitHub Pages.",
)

add_heading(doc, "2. Planeación del Sprint 4")
add_heading(doc, "2.1 Objetivo", level=2)
add_paragraph(
    doc,
    "El objetivo consistió en mejorar la exploración de la malla curricular mediante filtros accesibles y contadores dinámicos, sin modificar las reglas de prerrequisitos ni el avance almacenado por el usuario.",
)
add_heading(doc, "2.2 Historias de usuario comprometidas", level=2)
add_bullet(doc, "HU-09. Filtrar la malla por materias cursadas, disponibles y bloqueadas.")
add_bullet(doc, "HU-10. Consultar contadores que se actualizan cuando cambia el avance curricular.")
add_bullet(doc, "HU-11. Recibir una respuesta accesible cuando el filtro seleccionado no contiene resultados.")
add_heading(doc, "2.3 Organización del trabajo", level=2)
add_paragraph(
    doc,
    "Por tratarse de un prototipo académico individual, Pedro Pablo Rodríguez Gómez concentró las funciones de definición del alcance, desarrollo, revisión y prueba. El trabajo se organizó en análisis de las reglas existentes, implementación, verificación automatizada, compilación y publicación del incremento.",
)

page_break(doc)

# Página 3. Desarrollo
add_heading(doc, "3. Resultado del desarrollo del cuarto sprint")
add_paragraph(
    doc,
    "Las tres historias comprometidas fueron implementadas. El filtro activo se conserva como estado reactivo en App.vue; las reglas de clasificación y conteo se trasladaron a funciones puras en curriculum.js, y SemesterGroup.vue mantiene la responsabilidad de presentar únicamente las materias recibidas.",
)

add_table(
    doc,
    ["Historia", "Necesidad", "Resultado", "Evidencia"],
    [
        ("HU-09", "Explorar materias según su estado.", "Cuatro filtros: todas, cursadas, disponibles y bloqueadas.", "Interfaz publicada."),
        ("HU-10", "Conocer cuántas materias corresponden a cada estado.", "Contadores reactivos que cambian al marcar, desmarcar o restablecer el avance.", "Pruebas de conteo."),
        ("HU-11", "Comprender un resultado vacío y regresar a la vista general.", "Mensaje accesible y botón para mostrar nuevamente todas las materias.", "Estado sin resultados."),
    ],
    [0.75, 1.75, 2.8, 1.2],
)

add_heading(doc, "3.1 Cambios técnicos", level=2)
add_bullet(doc, "Se añadieron calcularEstadosMaterias, filtrarMateriasPorEstado y contarMateriasPorEstado como funciones reutilizables.")
add_bullet(doc, "Los botones de filtro comunican su selección mediante aria-pressed y funcionan con teclado.")
add_bullet(doc, "La distribución utiliza cuatro columnas en escritorio, dos en pantallas medianas y una en dispositivos pequeños.")
add_bullet(doc, "El README identifica el alcance del Sprint 4 y las historias HU-09, HU-10 y HU-11.")

add_heading(doc, "3.2 Criterios de aceptación alcanzados", level=2)
add_paragraph(
    doc,
    "La selección de un filtro cambia inmediatamente las materias y los semestres visibles; los contadores responden al estado real; los filtros no alteran localStorage; y la interfaz ofrece una vía clara para regresar a la vista completa cuando no existen coincidencias.",
)

page_break(doc)

# Página 4. Evidencias
add_heading(doc, "4. Verificación y evidencias")
add_paragraph(
    doc,
    "El incremento fue verificado antes de su publicación. Las reglas curriculares y los nuevos filtros se evaluaron mediante pruebas automatizadas; posteriormente se generó la compilación de producción y se confirmó el despliegue realizado por GitHub Actions.",
)

add_table(
    doc,
    ["Verificación", "Resultado", "Evidencia obtenida"],
    [
        ("Pruebas automatizadas", "10 de 10 aprobadas", "Clasificación, filtros, contadores, prerrequisitos y estadísticas."),
        ("Compilación", "Correcta", "Vite transformó 14 módulos y generó la carpeta dist."),
        ("Commit", "62159f6 en main", "Incremento del Sprint 4 registrado en GitHub."),
        ("GitHub Actions", "Ejecución exitosa", "Build y deploy concluyeron en 30 segundos."),
        ("Sitio público", "Actualizado", "Los cuatro filtros están disponibles en GitHub Pages."),
    ],
    [1.45, 1.45, 3.6],
)

add_heading(doc, "4.1 Enlaces del incremento", level=2)
links = [
    ("Aplicación publicada: ", "GitHub Pages", "https://pablovirtual.github.io/Prototipo-visualizador-de-avance/"),
    ("Código fuente: ", "Repositorio en GitHub", "https://github.com/pablovirtual/Prototipo-visualizador-de-avance"),
    ("Cierre del Sprint 4: ", "Commit 62159f6", "https://github.com/pablovirtual/Prototipo-visualizador-de-avance/commit/62159f64109accb5bcfca2949af18e4a43e48525"),
    ("Despliegue automático: ", "Ejecución 39002586849", "https://github.com/pablovirtual/Prototipo-visualizador-de-avance/actions/runs/39002586849"),
]
for lead, label, url in links:
    paragraph = add_paragraph(doc, after=5, line=1.2)
    add_hyperlink(paragraph, label, url, bold_lead=lead)

add_heading(doc, "4.2 Estado actual de la aplicación", level=2)
add_paragraph(
    doc,
    "El sitio mantiene el progreso, los créditos, la disponibilidad por prerrequisitos y la persistencia local desarrollados en los sprints anteriores. El filtro constituye una capa de consulta: no cambia el conjunto de materias cursadas ni modifica el JSON del plan de prueba.",
)

page_break(doc)

# Página 5. Retrospectiva y próximos pasos
add_heading(doc, "5. Resultado de la reunión retrospectiva")
add_heading(doc, "5.1 Aspectos favorables", level=2)
add_paragraph(
    doc,
    "La reutilización del mapa de estados permitió construir los filtros sin duplicar la lógica de prerrequisitos. La extracción de funciones puras facilitó las pruebas y mantuvo App.vue como coordinador del estado. La compilación y el despliegue automático ofrecieron una verificación inmediata del incremento publicado.",
)

add_heading(doc, "5.2 Dificultades identificadas", level=2)
add_paragraph(
    doc,
    "El principal reto consistió en diferenciar una materia disponible de una materia ya cursada, debido a que ambas cumplen sus prerrequisitos. También fue necesario mantener contadores coherentes después de cambios encadenados y adaptar los controles para pantallas pequeñas sin perder claridad ni acceso mediante teclado.",
)

add_heading(doc, "5.3 Lecciones aprendidas", level=2)
add_paragraph(
    doc,
    "La retrospectiva confirmó que las reglas de negocio resultan más sencillas de mantener cuando se expresan como funciones independientes de la interfaz. Asimismo, los criterios de aceptación deben contemplar resultados vacíos, navegación con teclado y comportamiento móvil desde el inicio del sprint.",
)

add_heading(doc, "6. Modificaciones propuestas para el Sprint 5")
add_paragraph(
    doc,
    "Para el siguiente incremento se proponen funcionalidades que mejoren la portabilidad del avance y amplíen la validación de la experiencia actual:",
)
add_bullet(doc, "Exportar el avance curricular como archivo JSON para crear una copia de respaldo.")
add_bullet(doc, "Importar un respaldo válido, ignorando IDs inexistentes y datos académicamente inconsistentes.")
add_bullet(doc, "Incorporar una búsqueda por nombre o identificador que funcione junto con los filtros de estado.")
add_bullet(doc, "Añadir pruebas de componentes y recorridos completos para filtros, persistencia, importación y exportación.")
add_bullet(doc, "Revisar accesibilidad, presentación móvil y mantenimiento del flujo de GitHub Actions.")

page_break(doc)

# Página 6. Conclusiones y referencias
add_heading(doc, "7. Conclusiones")
add_paragraph(
    doc,
    "El cuarto sprint cumplió el objetivo de mejorar la consulta de la malla curricular. Las historias HU-09, HU-10 y HU-11 se integraron en un incremento funcional que clasifica las materias, presenta conteos actualizados y comunica de forma comprensible la ausencia de resultados. Estas capacidades complementan la gestión de prerrequisitos desarrollada previamente.",
)
add_paragraph(
    doc,
    "Desde la perspectiva del estudiante, los filtros reducen el esfuerzo necesario para localizar asignaturas cursadas, disponibles o bloqueadas. La aplicación conserva las métricas y la persistencia local, mientras que los controles accesibles y la adaptación móvil amplían las condiciones de uso del prototipo. El resultado continúa siendo una herramienta de orientación y no sustituye la información académica oficial.",
)
add_paragraph(
    doc,
    "La verificación técnica respalda el avance alcanzado: las diez pruebas automatizadas concluyeron correctamente, la compilación de producción fue satisfactoria y GitHub Actions publicó el commit de cierre en GitHub Pages. El siguiente sprint deberá fortalecer la portabilidad del avance y las pruebas de interacción para continuar la evolución incremental del proyecto.",
)

add_heading(doc, "8. Referencias")
references = [
    "Atlassian. (s. f.). Métricas de Scrum. https://www.atlassian.com/es/agile/scrum/scrum-metrics",
    "Schwaber, K., & Sutherland, J. (2020). The Scrum Guide. Scrum Guides. https://scrumguides.org/scrum-guide.html",
    "Vue.js. (s. f.). Introduction. https://vuejs.org/guide/introduction.html",
    "MDN Web Docs. (s. f.). ARIA: aria-pressed attribute. https://developer.mozilla.org/en-US/docs/Web/Accessibility/ARIA/Reference/Attributes/aria-pressed",
]
for reference in references:
    paragraph = add_paragraph(doc, reference, after=5, line=1.15)
    paragraph.paragraph_format.left_indent = Inches(0.3)
    paragraph.paragraph_format.first_line_indent = Inches(-0.3)

doc.save(FINAL)
print(FINAL)
